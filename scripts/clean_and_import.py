#!/usr/bin/env python3
"""
Step 1: Hapus record Airtable yang No WA-nya tidak valid
Step 2: Import MASTER_DATABASE_LENGKAP_4756_KONTAK.xlsx (dedup by No WA)
"""

import re, time, json, urllib.request, urllib.parse
import openpyxl

TOKEN    = 'AIRTABLE_PAT_ISI_DI_ENV'
BASE_ID  = 'apprX0HpneI5lcqOP'
TABLE_ID = 'tbl7HTgx2LtFt6j9V'
BASE_URL = f'https://api.airtable.com/v0/{BASE_ID}/{TABLE_ID}'
HEADERS  = {'Authorization': f'Bearer {TOKEN}', 'Content-Type': 'application/json'}

MASTER_FILE = '/Users/imamzoel/Documents/nagara-digital-indonesia/Data/MASTER_DATABASE_LENGKAP_4756_KONTAK.xlsx'

CATEGORY_MAP = {
    'apartement': 'Kost & Apartemen',
    'coffee shop': 'Coffee Shop',
    'gym': 'Olahraga & Fitness',
    'hotel': 'Hotel & Penginapan',
    'klinik': 'Klinik & Kesehatan',
    'kost': 'Kost & Apartemen',
    'paud': 'Pendidikan',
    'perkantoran': 'Perkantoran',
    'restoran': 'Restoran',
    'salon dan barber': 'Salon & Barber',
    'sd': 'Pendidikan',
    'sma': 'Pendidikan',
    'smp': 'Pendidikan',
    'univ': 'Pendidikan',
    'spa & massage': 'Spa & Massage',
    'resto extra': 'Restoran',
}

UMROH_SHEETS = {'umroh-semua','bekasi','bdg','blitar','bogor','cianjur','cirebon',
                'depok','garut','indramayu','kuningan','lampung','majalengka',
                'palembang','purwakarta','semarang','solo','subang','sukabumi',
                'sumedang','tasik','yogyakarta'}

def normalize_wa(num):
    if not num:
        return ''
    n = re.sub(r'\D', '', str(num))
    if n.startswith('0'):
        n = '62' + n[1:]
    if not n.startswith('62'):
        n = '62' + n
    return n

def is_valid_wa(num):
    if not num:
        return False
    n = re.sub(r'\D', '', str(num))
    if n.startswith('0'):
        n = '62' + n[1:]
    if not n.startswith('62'):
        return False
    return 10 <= len(n) <= 15

def airtable_request(method, path='', data=None, params=None):
    url = BASE_URL + path
    if params:
        url += '?' + urllib.parse.urlencode(params)
    body = json.dumps(data).encode() if data else None
    req = urllib.request.Request(url, data=body, headers=HEADERS, method=method)
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read())

def fetch_all_records():
    print('Fetching semua record dari Airtable...')
    records = []
    offset = None
    while True:
        qs = 'fields[]=Nama&fields[]=No%20WA&pageSize=100'
        if offset:
            qs += '&offset=' + urllib.parse.quote(offset)
        url = BASE_URL + '?' + qs
        req = urllib.request.Request(url, headers=HEADERS, method='GET')
        with urllib.request.urlopen(req) as r:
            data = json.loads(r.read())
        records.extend(data.get('records', []))
        print(f'  fetched {len(records)}...', end='\r')
        offset = data.get('offset')
        if not offset:
            break
        time.sleep(0.2)
    print(f'\nTotal fetched: {len(records)}')
    return records

def delete_invalid(records):
    invalid_ids = []
    for r in records:
        wa = r['fields'].get('No WA', '')
        if not is_valid_wa(wa):
            invalid_ids.append(r['id'])

    print(f'\nRecord dengan No WA tidak valid: {len(invalid_ids)}')
    if not invalid_ids:
        print('Tidak ada yang perlu dihapus.')
        return set()

    deleted = 0
    for i in range(0, len(invalid_ids), 10):
        batch = invalid_ids[i:i+10]
        params = '&'.join(f'records[]={rid}' for rid in batch)
        url = BASE_URL + '?' + params
        req = urllib.request.Request(url, headers=HEADERS, method='DELETE')
        urllib.request.urlopen(req)
        deleted += len(batch)
        print(f'  deleted {deleted}/{len(invalid_ids)}...', end='\r')
        time.sleep(0.3)

    print(f'\nSelesai hapus {deleted} record invalid.')
    return set(invalid_ids)

def get_existing_wa(records):
    existing = set()
    for r in records:
        wa = r['fields'].get('No WA', '')
        if wa:
            existing.add(normalize_wa(wa))
    return existing

def read_master_xlsx(existing_wa):
    print(f'\nMembaca {MASTER_FILE}...')
    wb = openpyxl.load_workbook(MASTER_FILE, read_only=True, data_only=True)
    new_records = []
    seen = set(existing_wa)

    for sheet_name in wb.sheetnames:
        if sheet_name.startswith('📊'):
            continue
        clean_name = re.sub(r'[^\w\s&-]', '', sheet_name).strip().lower()

        is_umroh = any(u in clean_name for u in UMROH_SHEETS)
        if is_umroh:
            kategori = 'Travel Umroh & Haji'
        else:
            kategori = CATEGORY_MAP.get(clean_name, 'Lainnya')

        ws = wb[sheet_name]
        header = None
        count = 0
        for row in ws.iter_rows(values_only=True):
            if header is None:
                # Find header row
                row_lower = [str(c).lower().strip() if c else '' for c in row]
                if 'nama' in row_lower or 'nomor' in row_lower:
                    header = row_lower
                continue

            if not any(row):
                continue

            def get(field_fragments):
                for frag in field_fragments:
                    for i, h in enumerate(header):
                        if frag in h and i < len(row):
                            return str(row[i]).strip() if row[i] is not None else ''
                return ''

            nama = get(['nama'])
            wa_raw = get(['whatsapp', 'nomor', 'no wa', 'no. wa', 'phone'])
            alamat = get(['alamat', 'address'])
            website = get(['website', 'web'])

            if not nama or nama.startswith('='):
                continue

            wa = normalize_wa(wa_raw)
            if not is_valid_wa(wa):
                continue
            if wa in seen:
                continue
            seen.add(wa)

            # Extract kota dari alamat (last meaningful part)
            kota = ''
            if alamat:
                parts = [p.strip() for p in alamat.split(',') if p.strip()]
                if len(parts) >= 2:
                    kota = parts[-2] if len(parts) > 2 else parts[-1]
                    kota = kota[:50]

            catatan = f'Website: {website}' if website else ''

            new_records.append({
                'fields': {
                    'Nama': nama[:100],
                    'No WA': wa,
                    'Kota': kota,
                    'Kategori Blast': kategori,
                    'Catatan': catatan,
                }
            })
            count += 1

        if count:
            print(f'  {sheet_name}: +{count} kontak baru')

    wb.close()
    print(f'\nTotal kontak baru dari MASTER: {len(new_records)}')
    return new_records

def import_records(records):
    if not records:
        print('Tidak ada record baru untuk diimport.')
        return

    imported = 0
    errors = 0
    for i in range(0, len(records), 10):
        batch = records[i:i+10]
        try:
            airtable_request('POST', data={'records': batch, 'typecast': True})
            imported += len(batch)
        except Exception as e:
            errors += 1
            print(f'\nError batch {i//10+1}: {e}')
        print(f'\r[{imported}/{len(records)}] imported...', end='')
        time.sleep(0.3)

    print(f'\n\nSelesai! Imported: {imported}, Error: {errors}')

if __name__ == '__main__':
    # Step 1: bersihkan invalid
    print('='*50)
    print('STEP 1: BERSIHKAN NO WA TIDAK VALID')
    print('='*50)
    all_records = fetch_all_records()
    existing_wa = get_existing_wa(all_records)
    delete_invalid(all_records)

    # Refresh existing WA list (remove deleted ones)
    valid_wa = {normalize_wa(r['fields'].get('No WA',''))
                for r in all_records
                if is_valid_wa(r['fields'].get('No WA',''))}

    # Step 2: import MASTER
    print('\n' + '='*50)
    print('STEP 2: IMPORT MASTER DATABASE XLSX')
    print('='*50)
    new_records = read_master_xlsx(valid_wa)
    import_records(new_records)

    print('\nDone!')
